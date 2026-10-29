"""Розбір файлів з лідами: CSV, вставлені з Excel рядки або файл .xlsx.

Звідки беруться такі файли: вивантаження сканера бейджів на стенді,
список контактів із застосунку конференції, перелік від організатора
або таблиця, куди команда переписала візитки.
"""
import csv
import io

# Слова в заголовку таблиці, за якими впізнаємо стовпці (порівнюємо в нижньому регістрі)
HEADERS = {
    "full_name": ("ім'я та прізвище", "ім’я та прізвище", "піб", "full name", "name", "ім'я", "ім’я", "імя"),
    "first_name": ("first name", "firstname", "first"),
    "last_name": ("last name", "lastname", "surname", "прізвище"),
    "company": ("компанія", "company", "organization", "organisation", "організація"),
    "position": ("посада", "position", "job title", "title", "role"),
    "email": ("e-mail", "email", "пошта", "mail"),
    "phone": ("телефон", "phone", "mobile", "tel"),
    "linkedin_url": ("linkedin",),
    "intent": ("зацікавленість", "interest", "intent", "rating", "temperature"),
    "details": ("нотатка", "коментар", "notes", "note", "comment", "comments"),
}
# Порядок стовпців, якщо заголовка немає
DEFAULT_ORDER = ["full_name", "company", "position", "email", "phone", "linkedin_url", "intent", "details"]

INTENT_WORDS = {"3": 3, "гарячий": 3, "hot": 3, "high": 3, "a": 3,
                "2": 2, "теплий": 2, "warm": 2, "medium": 2, "b": 2,
                "1": 1, "холодний": 1, "cold": 1, "low": 1, "c": 1}


def rows_from_text(text):
    """Рядки з CSV або вставленої з Excel таблиці; розділювач визначається сам."""
    text = text.lstrip("﻿")
    try:
        dialect = csv.Sniffer().sniff(text[:2000], delimiters=";,\t")
    except csv.Error:
        dialect = csv.excel
    return [row for row in csv.reader(io.StringIO(text), dialect)]


def rows_from_xlsx(data):
    """Рядки першого аркуша файлу Excel. Потрібен пакет openpyxl."""
    from openpyxl import load_workbook
    sheet = load_workbook(io.BytesIO(data), read_only=True, data_only=True).active
    return [["" if v is None else str(v) for v in row] for row in sheet.iter_rows(values_only=True)]


def _match_header(cells):
    """Зіставлення стовпців за заголовком. Повертає None, якщо перший рядок не схожий на заголовок."""
    mapping = {}
    for i, cell in enumerate(cells):
        name = cell.strip().lower()
        for field, words in HEADERS.items():
            if field not in mapping.values() and any(name == w or (len(w) > 4 and w in name) for w in words):
                mapping[i] = field
                break
    # «Ім'я» поруч із «Прізвище» — це лише ім'я
    if "last_name" in mapping.values() and "first_name" not in mapping.values():
        mapping = {i: ("first_name" if f == "full_name" else f) for i, f in mapping.items()}
    has_name = "full_name" in mapping.values() or "first_name" in mapping.values()
    return mapping if has_name and len(mapping) >= 2 else None


def parse_leads(rows):
    """Перетворює рядки таблиці на ліди. Повертає (ліди, кількість пропущених рядків).

    Обов'язкове лише ім'я. Зацікавленість можна вказати словом (гарячий, теплий,
    холодний) або числом 3, 2, 1; якщо її немає, лід вважається теплим.
    """
    rows = [[str(c).strip() for c in r] for r in rows if any(str(c).strip() for c in r)]
    if not rows:
        return [], 0
    mapping = _match_header(rows[0])
    if mapping:
        rows = rows[1:]
    else:
        mapping = dict(enumerate(DEFAULT_ORDER))
    leads, skipped = [], 0
    for cells in rows:
        item = {field: cells[i] for i, field in mapping.items() if i < len(cells) and cells[i]}
        if not item.get("full_name"):
            item["full_name"] = " ".join(p for p in (item.get("first_name"), item.get("last_name")) if p)
        if not item.get("full_name"):
            skipped += 1
            continue
        contact = item.get("email", "")
        # Якщо в стовпці e-mail опинився LinkedIn або телефон, розкладаємо по місцях
        if contact and "@" not in contact:
            item.pop("email")
            if "linkedin.com" in contact:
                item.setdefault("linkedin_url", contact)
            else:
                item.setdefault("phone", contact)
        intent = INTENT_WORDS.get(str(item.get("intent", "")).strip().lower(), 2)
        leads.append({
            "full_name": item["full_name"][:200],
            "company": (item.get("company") or "")[:200] or None,
            "position": (item.get("position") or "")[:200] or None,
            "email": item.get("email"),
            "phone": item.get("phone"),
            "linkedin_url": item.get("linkedin_url"),
            "intent": intent,
            "details": item.get("details"),
            "channel": "email" if item.get("email") else ("linkedin" if item.get("linkedin_url")
                                                          else ("phone" if item.get("phone") else None)),
        })
    return leads, skipped
